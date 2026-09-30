import re
import uuid

import docx
import pdfplumber
from docx.document import Document as _Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.models import SpecDocument, SpecSection

HEADING_STYLES = {"Heading 1", "Heading 2"}
HEADING_NUMBER_RE = re.compile(r"^\s*(\d+(?:\.\d+)*)\.?\s+(.*)$")


# -----------------------------
# DOCX: walk paragraphs and tables in true document order
# -----------------------------
def _iter_block_items(parent):
    """
    Yield each paragraph and table in the order they actually appear in the
    document body. python-docx's `doc.paragraphs` / `doc.tables` give you
    each kind separately and lose that ordering, which is why the original
    parser never saw table content at all.
    """
    parent_elm = parent.element.body if isinstance(parent, _Document) else parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _table_to_text(table) -> str:
    """
    Render a table using its real cell boundaries (not flowed text), so
    values stay paired with the correct column even in irregular layouts.
    """
    rows = []
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        rows.append(" | ".join(cells))
    return "[TABLE]\n" + "\n".join(rows)


def parse_docx(path: str) -> SpecDocument:
    doc = docx.Document(path)
    sections = []
    warnings = []

    current_number = None
    current_heading = None
    current_text_parts = []
    fallback_counter = 1

    def flush():
        nonlocal fallback_counter
        text = "\n\n".join(p for p in current_text_parts if p.strip())
        if current_heading is not None and text.strip():
            section_id = f"SEC-{current_number}" if current_number else f"SEC-{fallback_counter}"
            if not current_number:
                fallback_counter += 1
            sections.append(
                SpecSection(
                    section_id=section_id,
                    heading=current_heading,
                    text=text,
                    page=None,
                )
            )

    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            style_name = block.style.name if block.style is not None else ""
            if style_name in HEADING_STYLES and block.text.strip():
                flush()
                current_text_parts = []
                match = HEADING_NUMBER_RE.match(block.text.strip())
                if match:
                    current_number, title = match.group(1), match.group(2)
                    current_heading = f"{current_number} {title}"
                else:
                    current_number = None
                    current_heading = block.text.strip()
            elif block.text.strip():
                current_text_parts.append(block.text.strip())
        elif isinstance(block, Table):
            current_text_parts.append(_table_to_text(block))

    flush()

    if not sections:
        warnings.append("No headings detected — falling back to no sections.")

    return SpecDocument(
        doc_id=str(uuid.uuid4()),
        source_filename=path,
        source_format="docx",
        title="PWCM Specification",
        sections=sections,
        ingestion_warnings=warnings,
    )


# -----------------------------
# PDF: best-effort — no reliable heading/table structure available,
# so we fall back to one section per page and flag it explicitly.
# -----------------------------
def parse_pdf(path: str) -> SpecDocument:
    sections = []
    section_counter = 1

    with pdfplumber.open(path) as pdf:
        for page_num, page in enumerate(pdf.pages):
            parts = []
            text = page.extract_text()
            if text:
                parts.append(text)
            for table in page.extract_tables():
                rows = [" | ".join(cell or "" for cell in row) for row in table]
                parts.append("[TABLE]\n" + "\n".join(rows))
            if parts:
                sections.append(
                    SpecSection(
                        section_id=f"SEC-p{page_num + 1}",
                        heading=f"Page {page_num + 1}",
                        text="\n\n".join(parts),
                        page=page_num + 1,
                    )
                )
                section_counter += 1

    return SpecDocument(
        doc_id=str(uuid.uuid4()),
        source_filename=path,
        source_format="pdf",
        title="PWCM Specification",
        sections=sections,
        ingestion_warnings=[
            "PDF ingestion is page-granular, not heading-granular — "
            "prefer the .docx source when both are available."
        ],
    )


def parse_spec(path: str) -> SpecDocument:
    if path.endswith(".docx"):
        return parse_docx(path)
    if path.endswith(".pdf"):
        return parse_pdf(path)
    raise ValueError("Unsupported format")
